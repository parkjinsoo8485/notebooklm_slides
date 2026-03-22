from pathlib import Path
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

W, H = 1920, 1080
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5

slides = [
    {
        "type": "cover",
        "scene": "봄 아침, 작은 마을과 무지개 숲이 멀리 보이는 따뜻한 수채화 풍경",
        "text": "웃음별을 찾아서",
        "colors": ((255, 214, 170), (255, 244, 214), (189, 225, 255)),
    },
    {
        "type": "story",
        "scene": "봄 아침, 작은 마을의 빨간 지붕 집 앞. 강아지 콩이가 하품하며 눈을 뜬다.",
        "text": "콩이는 오늘이 그냥 평범한 날일 줄 알았어요.\n그런데 문 앞에 반짝이는 금빛 편지 한 장이 놓여 있었어요.",
        "colors": ((255, 226, 188), (255, 247, 226), (199, 229, 255)),
    },
    {
        "type": "story",
        "scene": "콩이가 편지를 입에 물고 주인 소녀 하나에게 달려간다.",
        "text": "편지에는 이렇게 쓰여 있었어요.\n“무지개 숲의 웃음별이 사라졌어요. 도와주세요!”\n하나와 콩이는 서로를 보고 힘차게 고개를 끄덕였어요.",
        "colors": ((255, 220, 180), (255, 240, 210), (210, 233, 255)),
    },
    {
        "type": "story",
        "scene": "하나와 콩이가 배낭을 메고 마을 길을 출발한다.",
        "text": "배낭에는 물병, 쿠키, 손전등, 그리고 용기 한 스푼이 들어 있었어요.\n둘은 “우리가 찾아올게!” 하고 외치며 길을 나섰어요.",
        "colors": ((255, 210, 170), (255, 238, 205), (195, 226, 250)),
    },
    {
        "type": "story",
        "scene": "무지개 숲 입구, 나무들이 알록달록 빛나고 새들이 노래한다.",
        "text": "숲은 예뻤지만 길이 여러 갈래로 나뉘어 있었어요.\n그때 느릿느릿 거북이 토리가 나타나\n“서두르지 말고 바람 소리를 들어봐”라고 말했어요.",
        "colors": ((199, 236, 190), (236, 251, 220), (179, 220, 245)),
    },
    {
        "type": "story",
        "scene": "세 친구가 눈을 감고 바람 소리를 듣는다.",
        "text": "바람은 오른쪽 길에서 은은한 종소리를 실어 왔어요.\n하나와 콩이, 토리는 종소리를 따라 조심조심 걸어갔어요.",
        "colors": ((200, 234, 210), (238, 252, 232), (175, 215, 245)),
    },
    {
        "type": "story",
        "scene": "앞을 가로막은 출렁출렁 젤리 다리.",
        "text": "다리는 밟을 때마다 통통 튀어서 건너기 어려웠어요.\n콩이는 겁이 났지만 하나가 손을 내밀며 말했어요.\n“천천히, 같이 가면 돼!”",
        "colors": ((222, 206, 255), (244, 236, 255), (195, 225, 255)),
    },
    {
        "type": "story",
        "scene": "다리 건너편, 울고 있는 아기 구름 몽실.",
        "text": "몽실은 “내가 재채기하다가 웃음별을 날려 버렸어…” 하며 훌쩍였어요.\n모두는 혼내지 않고 “괜찮아, 같이 찾자!” 하고 몽실을 안아줬어요.",
        "colors": ((206, 224, 255), (236, 245, 255), (206, 236, 255)),
    },
    {
        "type": "story",
        "scene": "반짝이는 동굴 안, 높은 바위 틈에 걸린 웃음별을 발견한 순간.",
        "text": "웃음별은 손이 닿지 않는 곳에 있었어요.",
        "colors": ((140, 160, 210), (210, 225, 255), (150, 190, 235)),
    },
    {
        "type": "story",
        "scene": "토리가 받침이 되고 몽실이 몸을 띄우고, 하나 위로 콩이가 폴짝 뛰어 별을 잡는 장면.",
        "text": "토리는 받침이 되고, 몽실은 몸을 둥실 띄우고, 하나는 그 위에 올라섰어요.\n마지막으로 콩이가 폴짝 뛰어 웃음별을 톡! 잡았어요.",
        "colors": ((164, 178, 230), (224, 233, 255), (172, 205, 245)),
    },
    {
        "type": "story",
        "scene": "동굴 밖 하늘, 웃음별이 다시 떠오르며 숲 전체를 비춘다.",
        "text": "별이 번쩍하자 숲 전체에 웃음소리가 퍼졌어요.\n새들도 짹짹, 다람쥐도 까르르, 몽실도 킥킥 웃기 시작했어요.",
        "colors": ((255, 211, 150), (255, 240, 202), (173, 216, 255)),
    },
    {
        "type": "story",
        "scene": "무지개 숲 광장, 동물 친구들의 작은 축하 잔치.",
        "text": "모두가 쿠키를 나눠 먹으며 감사 인사를 했어요.\n하나는 말했어요.\n“우리가 해낸 건 힘이 세서가 아니라, 함께했기 때문이야.”",
        "colors": ((255, 222, 188), (255, 245, 218), (191, 231, 255)),
    },
    {
        "type": "story",
        "scene": "해질녘 마을로 돌아오는 길, 노을 아래 나란히 걷는 친구들.",
        "text": "콩이는 오늘이 세상에서 가장 멋진 날이라고 생각했어요.\n겁이 나도 손을 잡으면 용기가 생긴다는 걸 배웠거든요.",
        "colors": ((255, 182, 146), (255, 220, 186), (188, 210, 255)),
    },
    {
        "type": "story",
        "scene": "밤하늘 아래 집 앞, 하나와 콩이가 별을 올려다본다.",
        "text": "하늘 한가운데서 웃음별이 반짝이며 윙크했어요.",
        "colors": ((80, 102, 170), (163, 190, 245), (104, 140, 220)),
    },
    {
        "type": "ending",
        "scene": "같은 밤, 웃음별 클로즈업과 하나·콩이의 뒷모습. 따뜻하고 조용한 엔딩 무드.",
        "text": "하나가 속삭였어요.\n“내일도 누군가 도움이 필요하면, 우리 또 출동하자!”",
        "colors": ((70, 90, 155), (145, 172, 230), (95, 125, 200)),
    },
]

out_dir = Path("slides")
asset_dir = out_dir / "assets"
out_dir.mkdir(exist_ok=True)
asset_dir.mkdir(exist_ok=True)


def watercolor_bg(path: Path, c1, c2, c3, seed):
    img = Image.new("RGB", (W, H), c1)
    draw = ImageDraw.Draw(img, "RGBA")

    # Soft vertical blend bands
    for y in range(H):
        t = y / (H - 1)
        r = int(c1[0] * (1 - t) + c2[0] * t)
        g = int(c1[1] * (1 - t) + c2[1] * t)
        b = int(c1[2] * (1 - t) + c2[2] * t)
        draw.line((0, y, W, y), fill=(r, g, b, 255))

    # watercolor blobs
    import random
    rng = random.Random(seed)
    for _ in range(24):
        x = rng.randint(-200, W + 200)
        y = rng.randint(-150, H + 150)
        w = rng.randint(220, 620)
        h = rng.randint(140, 420)
        rr = int((c2[0] + c3[0]) / 2 + rng.randint(-18, 18))
        gg = int((c2[1] + c3[1]) / 2 + rng.randint(-18, 18))
        bb = int((c2[2] + c3[2]) / 2 + rng.randint(-18, 18))
        alpha = rng.randint(35, 70)
        draw.ellipse((x, y, x + w, y + h), fill=(max(0, min(255, rr)), max(0, min(255, gg)), max(0, min(255, bb)), alpha))

    # subtle stars/sparkles for kid-friendly feel
    for _ in range(70):
        x = rng.randint(0, W - 1)
        y = rng.randint(0, H - 1)
        rad = rng.randint(1, 3)
        draw.ellipse((x - rad, y - rad, x + rad, y + rad), fill=(255, 255, 255, rng.randint(45, 90)))

    img = img.filter(ImageFilter.GaussianBlur(radius=1.8))
    img.save(path)


for i, s in enumerate(slides, start=1):
    path = asset_dir / f"p{i:02}.png"
    watercolor_bg(path, *s["colors"], seed=100 + i)

prs = Presentation()
prs.slide_width = Inches(SLIDE_W_IN)
prs.slide_height = Inches(SLIDE_H_IN)
blank = prs.slide_layouts[6]

# fixed text box geometry
box_left = Inches(1.066)   # ~8%
box_top = Inches(5.72)     # bottom area fixed
box_width = Inches(11.2)   # ~84%
box_height = Inches(1.35)

for i, s in enumerate(slides, start=1):
    slide = prs.slides.add_slide(blank)

    # [1] Full-bleed single image
    slide.shapes.add_picture(str(asset_dir / f"p{i:02}.png"), Inches(0), Inches(0), width=Inches(SLIDE_W_IN), height=Inches(SLIDE_H_IN))

    # [3] Mandatory bottom box / cover label
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, box_left, box_top, box_width, box_height)
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(246, 237, 214)  # warm ivory
    fill.transparency = 8
    shape.line.fill.background()

    # Fixed corner style is represented by same rounded rectangle on all slides
    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(0.12)
    text_frame.margin_right = Inches(0.12)
    text_frame.margin_top = Inches(0.05)
    text_frame.margin_bottom = Inches(0.05)

    p = text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.4
    run = p.add_run()
    run.text = s["text"]

    font = run.font
    if s["type"] == "cover":
        font.name = "Noto Sans KR"
        font.bold = True
        font.size = Pt(64)
        font.color.rgb = RGBColor(56, 44, 33)
    else:
        font.name = "Noto Sans KR"
        font.bold = False
        font.size = Pt(28)
        font.color.rgb = RGBColor(43, 37, 30)

pptx_path = out_dir / "웃음별을-찾아서-15p.pptx"
prs.save(pptx_path)
print(f"Saved: {pptx_path}")
